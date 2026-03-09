"""PPT Generator — Creates professional PowerPoint presentations from sales data."""

import os
from datetime import datetime
from typing import Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from sqlalchemy.orm import Session

from app.models.stakeholder import Stakeholder
from app.models.company import Company
from app.models.campaign import OutreachCampaign, CampaignStakeholder
from app.models.outreach_message import OutreachMessage

# Colors
NAVY = RGBColor(0x00, 0x33, 0x66)
BLUE = RGBColor(0x00, 0x70, 0xF2)
GOLD = RGBColor(0xE8, 0xA3, 0x17)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MED_GRAY = RGBColor(0x66, 0x66, 0x66)

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "generated_ppts")


class PPTGenerator:

    def __init__(self):
        os.makedirs(OUTPUT_DIR, exist_ok=True)

    def _new_prs(self) -> Presentation:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        return prs

    def _title_slide(self, prs, title: str, subtitle: str):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = NAVY

        # Gold accent line
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.8), Inches(2), Pt(3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = GOLD
        shape.line.fill.background()

        # Title
        tb = slide.shapes.add_textbox(Inches(1), Inches(3.1), Inches(11), Inches(1.5))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(40), True, WHITE, "Calibri"

        # Subtitle
        tb2 = slide.shapes.add_textbox(Inches(1), Inches(4.6), Inches(11), Inches(1))
        p2 = tb2.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.font.size, p2.font.color.rgb, p2.font.name = Pt(18), GOLD, "Calibri"

        self._footer(slide, dark=True)

    def _content_slide(self, prs, title: str, blocks: list[dict]):
        """blocks: list of {'heading':str, 'bullets':[str], 'text':str, 'kvs':[(k,v)]}"""
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Title bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(0.75), Inches(0.2), Inches(11.5), Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(28), True, WHITE, "Calibri"

        top = Inches(1.4)
        for block in blocks:
            top = self._render_block(slide, Inches(0.75), top, Inches(11.5), block)

        self._footer(slide)

    def _two_col_slide(self, prs, title: str, left: dict, right: dict):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(0.75), Inches(0.2), Inches(11.5), Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(28), True, WHITE, "Calibri"

        self._render_block(slide, Inches(0.75), Inches(1.4), Inches(5.5), left)
        # Divider
        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.5), Inches(1.4), Pt(2), Inches(5.2))
        div.fill.solid()
        div.fill.fore_color.rgb = RGBColor(0xEB, 0xF5, 0xFB)
        div.line.fill.background()
        self._render_block(slide, Inches(6.9), Inches(1.4), Inches(5.5), right)
        self._footer(slide)

    def _render_block(self, slide, left, top, width, block: dict):
        if block.get("heading"):
            tb = slide.shapes.add_textbox(left, top, width, Inches(0.4))
            tb.text_frame.word_wrap = True
            p = tb.text_frame.paragraphs[0]
            p.text = block["heading"]
            p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(18), True, BLUE, "Calibri"
            top += Inches(0.45)

        if block.get("text"):
            tb = slide.shapes.add_textbox(left, top, width, Inches(0.8))
            tb.text_frame.word_wrap = True
            p = tb.text_frame.paragraphs[0]
            p.text = block["text"]
            p.font.size, p.font.color.rgb, p.font.name = Pt(14), DARK_GRAY, "Calibri"
            top += Inches(0.7)

        for kv in block.get("kvs", []):
            tb = slide.shapes.add_textbox(left, top, width, Inches(0.35))
            tb.text_frame.word_wrap = True
            p = tb.text_frame.paragraphs[0]
            r1 = p.add_run()
            r1.text = f"{kv[0]}: "
            r1.font.size, r1.font.bold, r1.font.color.rgb, r1.font.name = Pt(13), True, NAVY, "Calibri"
            r2 = p.add_run()
            r2.text = str(kv[1])
            r2.font.size, r2.font.color.rgb, r2.font.name = Pt(13), DARK_GRAY, "Calibri"
            top += Inches(0.32)

        for bullet in block.get("bullets", []):
            tb = slide.shapes.add_textbox(left + Inches(0.3), top, width - Inches(0.3), Inches(0.35))
            tb.text_frame.word_wrap = True
            p = tb.text_frame.paragraphs[0]
            p.text = f"\u2022  {bullet}"
            p.font.size, p.font.color.rgb, p.font.name = Pt(13), DARK_GRAY, "Calibri"
            top += Inches(0.3)

        return top + Inches(0.15)

    def _footer(self, slide, dark=False):
        tb = slide.shapes.add_textbox(Inches(0.75), Inches(7.0), Inches(11.5), Inches(0.3))
        p = tb.text_frame.paragraphs[0]
        p.text = "SAP Sales Outreach Engine  |  Confidential"
        p.font.size, p.font.italic = Pt(9), True
        p.font.color.rgb = WHITE if dark else MED_GRAY
        p.font.name = "Calibri"
        p.alignment = PP_ALIGN.RIGHT

    def _save(self, prs, prefix: str, entity_id=None) -> dict:
        suffix = f"_{entity_id}" if entity_id else ""
        filename = f"{prefix}{suffix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
        filepath = os.path.join(OUTPUT_DIR, filename)
        prs.save(filepath)
        return {"filename": filename, "filepath": filepath, "download_url": f"/api/presentations/download/{filename}", "slides_count": len(prs.slides), "type": prefix}

    # ── Generators ──────────────────────────────────────────────────

    async def generate_stakeholder_briefing(self, stakeholder: Stakeholder, company: Optional[Company], db: Session, title: Optional[str] = None, ai_provider: str = "ollama") -> dict:
        prs = self._new_prs()
        t = title or f"Stakeholder Briefing: {stakeholder.full_name}"
        sub = f"{stakeholder.job_title or 'Contact'} at {stakeholder.company_name or 'Unknown Company'}"
        self._title_slide(prs, t, sub)

        # Slide 2: Profile
        self._two_col_slide(prs, "Profile Overview",
            {"heading": "Professional Profile", "kvs": [
                ("Name", stakeholder.full_name),
                ("Title", stakeholder.job_title or "N/A"),
                ("Department", stakeholder.department or "N/A"),
                ("Seniority", (stakeholder.seniority_level or "N/A").replace("_", " ").title()),
                ("Location", stakeholder.linkedin_location or "N/A"),
                ("Decision Role", (stakeholder.decision_maker_type or "N/A").replace("_", " ").title()),
            ]},
            {"heading": "Background",
             "text": stakeholder.linkedin_summary or stakeholder.linkedin_headline or "No LinkedIn summary available.",
             "bullets": [f"Interest: {i}" for i in (stakeholder.interests or [])[:5]]},
        )

        # Slide 3: Need & Gap Analysis
        blocks = []
        if stakeholder.ai_need_analysis:
            blocks.append({"heading": "Need Analysis", "text": stakeholder.ai_need_analysis[:500]})
        if stakeholder.ai_gap_analysis:
            blocks.append({"heading": "Gap Analysis", "text": stakeholder.ai_gap_analysis[:500]})
        if stakeholder.pain_points:
            blocks.append({"heading": "Identified Pain Points", "bullets": stakeholder.pain_points[:6]})
        if not blocks:
            blocks.append({"text": "No analysis data available. Run stakeholder profiling first."})
        self._content_slide(prs, "Need & Gap Analysis", blocks)

        # Slide 4: SAP Solutions
        blocks = []
        if stakeholder.recommended_sap_solutions:
            blocks.append({"heading": "Recommended SAP Solutions", "bullets": stakeholder.recommended_sap_solutions[:6]})
        if stakeholder.current_sap_experience:
            blocks.append({"heading": "Current SAP Experience", "text": stakeholder.current_sap_experience})
        if company and company.current_sap_products:
            blocks.append({"heading": "Company SAP Products", "bullets": company.current_sap_products[:6]})
        if not blocks:
            blocks.append({"text": "No SAP solution recommendations yet. Run profiling to generate."})
        self._content_slide(prs, "SAP Solution Recommendations", blocks)

        # Slide 5: Engagement Strategy
        left = {"heading": "Conversation Starters"}
        if stakeholder.conversation_starters:
            left["bullets"] = stakeholder.conversation_starters[:4]
        else:
            left["text"] = "Run profiling to generate conversation starters."
        right = {"heading": "Engagement Details", "kvs": [
            ("Best Channel", (stakeholder.best_outreach_channel or "LinkedIn").replace("_", " ").title()),
            ("Engagement Score", f"{stakeholder.engagement_score or 0:.0f}/100"),
            ("Fit Score", f"{stakeholder.fit_score or 0:.0f}/100"),
            ("Buying Stage", (stakeholder.buying_stage or "Unknown").replace("_", " ").title()),
            ("Budget Authority", (stakeholder.budget_authority or "Unknown").title()),
        ]}
        self._two_col_slide(prs, "Engagement Strategy", left, right)

        # Slide 6: Talking Points
        blocks = []
        if stakeholder.content_topics:
            blocks.append({"heading": "Content Topics of Interest", "bullets": stakeholder.content_topics[:5]})
        if stakeholder.ai_profile_summary:
            blocks.append({"heading": "AI Profile Summary", "text": stakeholder.ai_profile_summary[:400]})
        if not blocks:
            blocks.append({"text": "Run stakeholder profiling to generate talking points."})
        self._content_slide(prs, "Key Talking Points", blocks)

        return self._save(prs, "stakeholder_briefing", stakeholder.id)

    async def generate_company_analysis(self, company: Company, db: Session, title: Optional[str] = None, ai_provider: str = "ollama") -> dict:
        prs = self._new_prs()
        t = title or f"Company Analysis: {company.name}"
        self._title_slide(prs, t, f"{company.industry or 'Industry Unknown'}  |  Tier {company.priority_tier or 3}")

        # Overview
        self._two_col_slide(prs, "Company Overview",
            {"heading": "Details", "kvs": [
                ("Company", company.name), ("Industry", company.industry or "N/A"),
                ("Sub-Industry", company.sub_industry or "N/A"), ("Employees", company.employee_count or "N/A"),
                ("Revenue", company.revenue_range or "N/A"), ("HQ", company.headquarters_location or "N/A"),
            ]},
            {"heading": "Classification", "kvs": [
                ("Priority Tier", f"Tier {company.priority_tier or 3}"),
                ("Digital Transformation", (company.digital_transformation_stage or "Unknown").replace("_", " ").title()),
                ("Regions", ", ".join(company.regions) if company.regions else "N/A"),
            ]},
        )

        # Tech Landscape
        blocks = [{"heading": "Current ERP & SAP Landscape", "kvs": [
            ("Current ERP", company.current_erp or "Unknown"),
            ("SAP Maturity", (company.sap_maturity_level or "Unknown").title()),
        ]}]
        if company.current_sap_products:
            blocks.append({"heading": "Current SAP Products", "bullets": company.current_sap_products})
        else:
            blocks.append({"text": "No SAP products currently deployed."})
        self._content_slide(prs, "Current Technology Landscape", blocks)

        # Pain Points
        blocks = []
        if company.known_pain_points:
            blocks.append({"heading": "Identified Pain Points", "bullets": company.known_pain_points[:8]})
        if company.ai_company_analysis:
            blocks.append({"heading": "AI Analysis", "text": company.ai_company_analysis[:500]})
        if not blocks:
            blocks.append({"text": "No pain points identified. Run company analysis to generate insights."})
        self._content_slide(prs, "Pain Points & Gaps", blocks)

        # Key Stakeholders
        shs = db.query(Stakeholder).filter(Stakeholder.company_id == company.id).limit(10).all()
        if shs:
            bullets = [f"{s.full_name} — {s.job_title or 'Unknown Role'}" + (f" (Fit: {s.fit_score:.0f})" if s.fit_score else "") for s in shs]
            self._content_slide(prs, "Key Stakeholders", [{"heading": f"Contacts at {company.name} ({len(shs)})", "bullets": bullets}])
        else:
            self._content_slide(prs, "Key Stakeholders", [{"text": f"No stakeholders found for {company.name}."}])

        # Next Steps
        self._content_slide(prs, "Recommended Approach & Next Steps", [{"heading": "Next Steps", "bullets": [
            "Complete stakeholder profiling for all key contacts",
            "Run AI analysis to identify SAP solution fits",
            "Generate personalized outreach for top-priority contacts",
            "Schedule discovery calls with decision makers",
            "Prepare tailored SAP solution demo",
        ]}])

        return self._save(prs, "company_analysis", company.id)

    async def generate_campaign_summary(self, campaign: OutreachCampaign, db: Session, title: Optional[str] = None, ai_provider: str = "ollama") -> dict:
        prs = self._new_prs()
        t = title or f"Campaign: {campaign.name}"
        self._title_slide(prs, t, f"Status: {(campaign.status or 'Draft').title()}  |  Channel: {(campaign.channel or 'LinkedIn').replace('_', ' ').title()}")

        sh_count = db.query(CampaignStakeholder).filter(CampaignStakeholder.campaign_id == campaign.id).count()
        msg_count = db.query(OutreachMessage).filter(OutreachMessage.campaign_id == campaign.id).count()

        # Overview
        left = {"heading": "Campaign Details", "kvs": [
            ("Name", campaign.name), ("Status", (campaign.status or "draft").title()),
            ("Channel", (campaign.channel or "linkedin_message").replace("_", " ").title()),
            ("Objective", campaign.objective or "N/A"),
        ]}
        right = {"heading": "Metrics", "kvs": [("Target Stakeholders", str(sh_count)), ("Messages Generated", str(msg_count))]}
        if campaign.sap_solution_focus:
            right["bullets"] = [f"SAP Focus: {s}" for s in campaign.sap_solution_focus]
        self._two_col_slide(prs, "Campaign Overview", left, right)

        # Stakeholder Breakdown
        sh_ids = [cs.stakeholder_id for cs in db.query(CampaignStakeholder).filter(CampaignStakeholder.campaign_id == campaign.id).all()]
        shs = db.query(Stakeholder).filter(Stakeholder.id.in_(sh_ids)).all() if sh_ids else []
        sen, dep = {}, {}
        for s in shs:
            k = (s.seniority_level or "unknown").replace("_", " ").title()
            sen[k] = sen.get(k, 0) + 1
            k2 = (s.department or "Unknown").title()
            dep[k2] = dep.get(k2, 0) + 1
        self._two_col_slide(prs, "Target Stakeholder Breakdown",
            {"heading": "By Seniority", "bullets": [f"{k}: {v}" for k, v in sorted(sen.items(), key=lambda x: -x[1])] or ["No data"]},
            {"heading": "By Department", "bullets": [f"{k}: {v}" for k, v in sorted(dep.items(), key=lambda x: -x[1])] or ["No data"]},
        )

        # SAP Focus
        blocks = []
        if campaign.sap_solution_focus:
            blocks.append({"heading": "SAP Solutions Being Positioned", "bullets": campaign.sap_solution_focus})
        if campaign.messaging_guidelines:
            blocks.append({"heading": "Messaging Guidelines", "text": campaign.messaging_guidelines[:400]})
        if not blocks:
            blocks.append({"text": "No specific SAP solutions targeted."})
        self._content_slide(prs, "SAP Solutions Focus", blocks)

        # Sample Messages
        msgs = db.query(OutreachMessage).filter(OutreachMessage.campaign_id == campaign.id).limit(2).all()
        blocks = [{"heading": "Message Stats", "kvs": [("Total Messages", str(msg_count))]}]
        for i, m in enumerate(msgs):
            preview = (m.body or "")[:200] + ("..." if len(m.body or "") > 200 else "")
            blocks.append({"heading": f"Sample Message {i+1}", "text": preview})
        if not msgs:
            blocks.append({"text": "No messages generated yet."})
        self._content_slide(prs, "Messaging Stats & Samples", blocks)

        # Next Steps
        self._content_slide(prs, "Next Steps", [{"heading": "Actions", "bullets": [
            "Review and approve generated messages",
            "Generate A/B variants for top-priority stakeholders",
            "Begin outreach sequence execution",
            "Track engagement and reply rates",
            "Iterate messaging based on responses",
        ]}])

        return self._save(prs, "campaign_summary", campaign.id)

    async def generate_outreach_strategy(self, db: Session, title: Optional[str] = None, ai_provider: str = "ollama") -> dict:
        prs = self._new_prs()
        total_sh = db.query(Stakeholder).count()
        total_co = db.query(Company).count()
        total_camp = db.query(OutreachCampaign).count()
        self._title_slide(prs, title or "SAP Outreach Strategy", f"{total_sh} Contacts  |  {total_co} Companies  |  {total_camp} Campaigns")

        # Segmentation
        shs = db.query(Stakeholder).all()
        sen, dep = {}, {}
        for s in shs:
            k = (s.seniority_level or "unknown").replace("_", " ").title()
            sen[k] = sen.get(k, 0) + 1
            k2 = (s.department or "Unknown").title()
            dep[k2] = dep.get(k2, 0) + 1
        self._two_col_slide(prs, "Target Audience Segmentation",
            {"heading": "By Seniority", "bullets": [f"{k}: {v}" for k, v in sorted(sen.items(), key=lambda x: -x[1])[:8]] or ["No data"]},
            {"heading": "By Department", "bullets": [f"{k}: {v}" for k, v in sorted(dep.items(), key=lambda x: -x[1])[:8]] or ["No data"]},
        )

        # Channel Strategy
        self._content_slide(prs, "Channel Strategy", [
            {"heading": "LinkedIn Outreach", "bullets": [
                "Connection requests with personalized notes (300 char limit)",
                "Follow-up InMails for non-connections",
                "Value-share messages with industry insights",
                "Best for: C-Suite, VP, Director-level contacts",
            ]},
            {"heading": "Email Outreach", "bullets": [
                "Personalized emails with subject line optimization",
                "Multi-touch sequences (3-5 touchpoints over 2-3 weeks)",
                "Best for: Manager-level contacts and follow-ups",
            ]},
        ])

        # Message Framework
        self._content_slide(prs, "Message Framework & Tone", [
            {"heading": "Message Types", "bullets": [
                "Initial Connect — Hook, Relevance, Soft Ask",
                "Follow Up — Callback, Value, Question",
                "Value Share — Context, Insight, Connection, Soft CTA",
                "Meeting Request — Recap, Specific Value, Clear Ask, Easy Out",
            ]},
            {"heading": "Tone Options", "bullets": [
                "Professional — Consultative, peer-level",
                "Casual — Warm, conversational",
                "Thought Leader — Insight-led, expert positioning",
            ]},
        ])

        # SAP Positioning
        sol_counts = {}
        for s in shs:
            for sol in (s.recommended_sap_solutions or []):
                sol_counts[sol] = sol_counts.get(sol, 0) + 1
        if sol_counts:
            bullets = [f"{k} ({v} stakeholders)" for k, v in sorted(sol_counts.items(), key=lambda x: -x[1])[:10]]
        else:
            bullets = ["SAP S/4HANA Cloud", "SAP Business AI", "SAP BTP", "SAP SuccessFactors", "SAP Analytics Cloud"]
        self._content_slide(prs, "SAP Solution Positioning", [{"heading": "Top Solutions", "bullets": bullets}])

        # Timeline
        self._content_slide(prs, "Timeline & Milestones", [{"heading": "Execution Plan", "bullets": [
            "Week 1-2: Import contacts, run AI profiling, segment audience",
            "Week 3: Generate personalized outreach and A/B variants",
            "Week 4-5: Execute campaigns across channels",
            "Week 6: Analyze response rates, iterate messaging",
            "Week 7-8: Follow-up sequences for engaged contacts",
        ]}])

        return self._save(prs, "outreach_strategy")


ppt_generator = PPTGenerator()
