"""
Fix remaining issues in the Toyota AU customised v2.0 PPTX:
1. Double 'Why Why' → 'Why'
2. Garbled 'weeks aheadonths ahead' text
3. 'freed/plant' → 'freed/depot' in ROI table
4. HiLux reference in demand forecasting subtitle
5. Additional Toyota AU polish
"""

from pptx import Presentation
from pptx.util import Emu

INPUT  = "backend/generated_ppts/NDBS_SAP_Supply_Chain_AI_v2.0.pptx"
OUTPUT = "backend/generated_ppts/NDBS_SAP_Supply_Chain_AI_v2.0.pptx"

prs = Presentation(INPUT)

fixes = 0
for slide_idx, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                orig = run.text

                # Fix double "Why Why"
                if "Why Why" in run.text:
                    run.text = run.text.replace("Why Why", "Why")

                # Fix garbled demand forecasting subtitle
                if "weeks aheadonths ahead" in run.text:
                    run.text = run.text.replace(
                        "Predict dealership demand for vehicles and parts weeks aheadonths ahead",
                        "Predict dealer demand for HiLux, RAV4 & parts across 275 AU dealers")
                if "weeks aheadonths" in run.text:
                    run.text = run.text.replace("weeks aheadonths ahead", "across Australia")

                # Fix freed/plant → freed/depot in ROI table
                if "freed/plant" in run.text:
                    run.text = run.text.replace("freed/plant", "freed/depot")

                # Fix "Optimisation" → keep consistent (already correct)

                # Fix remaining generic "your clients"
                if "your clients" in run.text.lower():
                    run.text = run.text.replace("your clients", "Toyota Australia")
                    run.text = run.text.replace("Your clients", "Toyota Australia")

                # Fix remaining "SAP clients"
                if "SAP clients" in run.text:
                    run.text = run.text.replace("SAP clients", "automotive OEMs")

                if run.text != orig:
                    fixes += 1
                    print(f"  S{slide_idx+1}: \"{orig[:60]}\" → \"{run.text[:60]}\"")

prs.save(OUTPUT)
print(f"\nFixes applied: {fixes}")
